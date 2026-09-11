# app/api/endpoints/management.py
import asyncio
import io
import csv
import re
import json
import secrets
from pathlib import Path
from typing import Optional, List
from fastapi import APIRouter, Depends, Query, HTTPException, status, UploadFile, File, Form, Request, Response
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, delete, update
from collections import defaultdict
from app.database.session import get_db
from app.database.models import Card, ReviewLog, Phrase, UserSession, DailySession, UserSetting, GenerationJob
from app.services.ai_gateway import parse_raw_text, regenerate_card_mnemonic, split_text_into_chunks
from app.services.generation_worker import is_deepseek_offpeak
from app.core.auth import get_current_user_id
from app.core.config import settings
from datetime import datetime, timedelta

router = APIRouter()

class ShareDeckIn(BaseModel):
    subject: str

class ConfigUpdate(BaseModel):
    daily_limit: int
    focus_mode_default: bool = False
    target_retention: float | None = 0.9
    assoc_preference: str | None = "acoustic"

class ImportIn(BaseModel):
    text: str
    subject: str = ""                    # Целевой предмет, выбранный человеком
    density: str = "medium"
    volume: str = "medium"
    priority: str = "balanced"
    assoc_preference: str = "acoustic"
    granularity_mode: str = "atomic"     # "atomic" | "single_deep" | "cheatsheet"
    custom_instruction: str = ""        # Свободные пожелания пользователя
    commit_now: bool = False            # False = вернуть в Песочницу (Staging)
    is_deferred: bool = False           # True = отправить в очередь Ночного Грайндера (-50% стоимости)

class PresetImportIn(BaseModel):
    preset_name: str
    commit_now: bool = False  # False = вернуть в Песочницу (Staging)

class CardStagingItem(BaseModel):
    text: str
    secondary_text: Optional[str] = ""
    translation: str
    example: Optional[str] = ""
    initial_difficulty_tier: Optional[str] = "medium"
    mnemonic: dict | str | None = None
    theme: Optional[str] = ""

class StagingCommitIn(BaseModel):
    subject: str
    theme: str
    cards: list[CardStagingItem]
    job_id: Optional[int] = None


class ManualCardIn(BaseModel):
    subject: str
    phrase_title: str = "Пользовательские карточки"
    text: str
    secondary_text: str = ""
    translation: str
    example: str = ""
    difficulty: float = 5.0
    mnemonic_keyword: str = ""
    mnemonic_cue: str = ""

class CardUpdateIn(BaseModel):
    text: str
    secondary_text: str = ""
    translation: str
    example: str = ""
    mnemonic_keyword: str = ""
    mnemonic_cue: str = ""

class CardMoveIn(BaseModel):
    target_subject: str

class BulkCardMoveIn(BaseModel):
    card_ids: list[int]
    target_subject: str

class BulkCardDeleteIn(BaseModel):
    card_ids: list[int]

class RegenerateMnemonicIn(BaseModel):
    preference: str = "visual"

class SubjectRenameIn(BaseModel):
    old_subject: str
    new_subject: str

class DailySessionIn(BaseModel):
    mental_effort: int
    association_utility: int
    perceived_retention: int
    session_duration: int

async def check_experiment_lock(current_user: str, db: AsyncSession):
    """Проверяет блокировку модификации колоды и настроек для участников научного эксперимента (Фаза 1)."""
    session_res = await db.execute(select(UserSession).filter(UserSession.user_id == current_user))
    user_sess = session_res.scalars().first()
    if user_sess and user_sess.is_experiment_participant and user_sess.experiment_phase == 1:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Действие заблокировано на период проведения научного эксперимента"
        )

# Вспомогательная функция для создания карточек в БД
async def save_cards_to_database(cards_data: list, subject_slug: str, phrase_title: str, user_id: str, db: AsyncSession):
    clean_sub = subject_slug.strip().lower() or "generic"
    clean_title = phrase_title.strip() or "Новый блок знаний"
    
    # Кэш тем (Phrases) для поддержки мульти-тематической кластеризации в одном пакете карточек
    phrase_cache = {}

    cards_created = 0
    now = datetime.utcnow()
    for c in cards_data:
        c_text = c.get("text", "") if isinstance(c, dict) else getattr(c, "text", "")
        c_trans = c.get("translation", "") if isinstance(c, dict) else getattr(c, "translation", "")
        if not c_text or not c_trans:
            continue

        c_theme = (c.get("theme", "") if isinstance(c, dict) else getattr(c, "theme", "") or "").strip() or clean_title

        if c_theme not in phrase_cache:
            phrase_res = await db.execute(
                select(Phrase).filter(Phrase.text == c_theme, Phrase.subject == clean_sub, Phrase.user_id == user_id)
            )
            phrase = phrase_res.scalar_one_or_none()
            if not phrase:
                phrase = Phrase(text=c_theme, subject=clean_sub, user_id=user_id)
                db.add(phrase)
                await db.flush()
            phrase_cache[c_theme] = phrase

        target_phrase = phrase_cache[c_theme]

        c_sec = c.get("secondary_text", "") if isinstance(c, dict) else getattr(c, "secondary_text", "")
        c_ex = c.get("example", "") if isinstance(c, dict) else getattr(c, "example", "")
        c_tier = c.get("initial_difficulty_tier", "medium") if isinstance(c, dict) else getattr(c, "initial_difficulty_tier", "medium")
        c_mnem = c.get("mnemonic", None) if isinstance(c, dict) else getattr(c, "mnemonic", None)

        difficulty = 5.5
        if c_tier == "easy":
            difficulty = 3.5
        elif c_tier == "hard":
            difficulty = 7.5

        stability = 1.0
        if c_mnem:
            if isinstance(c_mnem, dict) and c_mnem.get("keyword"):
                stability = 1.5
            elif isinstance(c_mnem, str) and c_mnem.strip():
                stability = 1.5

        c_type = (c.get("content_type") if isinstance(c, dict) else getattr(c, "content_type", None)) or ("cloze" if "{{c" in c_text else "text")

        card = Card(
            phrase_id=target_phrase.id,
            user_id=user_id,
            subject=clean_sub,
            text=c_text,
            secondary_text=c_sec,
            translation=c_trans,
            example=c_ex,
            difficulty=difficulty,
            stability=stability,
            state=0,
            mnemonic=c_mnem,
            content_type=c_type,
            next_review=now
        )
        db.add(card)
        cards_created += 1

    return cards_created, clean_sub, clean_title

async def append_or_sync_cards_to_database(
    cards_data: list,
    subject_slug: str,
    phrase_title: str,
    user_id: str,
    db: AsyncSession
) -> tuple[int, int, str, str]:
    """
    Дозагружает и синхронизирует карточки для конкретного пользователя:
    - Существующие карточки определяются по совпадению нормализованного текста вопроса (text.strip().lower()).
      Для них обновляются формулировки (translation, secondary_text, example, mnemonic),
      но полностью сохраняется когнитивный прогресс FSRS v4 (state, stability, difficulty, next_review, last_review, lapses, reps, has_seen_intro).
    - Новые карточки добавляются в базу со state=0 и next_review=now.
    Возвращает (cards_created, cards_updated, clean_sub, clean_title).
    """
    clean_sub = subject_slug.strip().lower() or "generic"
    clean_title = phrase_title.strip() or "Новый блок знаний"

    # Загружаем существующие карточки пользователя по данному предмету
    stmt_existing = select(Card).filter(Card.user_id == user_id, Card.subject == clean_sub)
    res_existing = await db.execute(stmt_existing)
    existing_cards = res_existing.scalars().all()
    existing_map = {c.text.strip().lower(): c for c in existing_cards if c.text}

    # Кэш тем (Phrases)
    stmt_phrases = select(Phrase).filter(Phrase.user_id == user_id, Phrase.subject == clean_sub)
    res_phrases = await db.execute(stmt_phrases)
    phrase_cache = {p.text.strip(): p for p in res_phrases.scalars().all() if p.text}

    cards_created = 0
    cards_updated = 0
    now = datetime.utcnow()

    for c in cards_data:
        c_text = (c.get("text", "") if isinstance(c, dict) else getattr(c, "text", "")) or ""
        c_trans = (c.get("translation", "") if isinstance(c, dict) else getattr(c, "translation", "")) or ""
        if not c_text.strip() or not c_trans.strip():
            continue

        c_text_clean = c_text.strip()
        c_key = c_text_clean.lower()
        c_sec = (c.get("secondary_text", "") if isinstance(c, dict) else getattr(c, "secondary_text", "")) or ""
        c_ex = (c.get("example", "") if isinstance(c, dict) else getattr(c, "example", "")) or ""
        c_tier = (c.get("initial_difficulty_tier", "medium") if isinstance(c, dict) else getattr(c, "initial_difficulty_tier", "medium"))
        c_mnem = c.get("mnemonic", None) if isinstance(c, dict) else getattr(c, "mnemonic", None)

        if c_key in existing_map:
            # Существующая карточка: обновляем только текстовые поля, сохраняя весь прогресс FSRS
            card = existing_map[c_key]
            card.translation = c_trans
            if c_sec:
                card.secondary_text = c_sec
            if c_ex:
                card.example = c_ex
            if c_mnem is not None:
                card.mnemonic = c_mnem
            cards_updated += 1
        else:
            # Новая карточка: привязываем к Phrase и добавляем в очередь
            c_theme = ((c.get("theme", "") if isinstance(c, dict) else getattr(c, "theme", "")) or "").strip() or clean_title

            if c_theme not in phrase_cache:
                phrase = Phrase(text=c_theme, subject=clean_sub, user_id=user_id)
                db.add(phrase)
                await db.flush()
                phrase_cache[c_theme] = phrase

            target_phrase = phrase_cache[c_theme]

            difficulty = 5.5
            if c_tier == "easy":
                difficulty = 3.5
            elif c_tier == "hard":
                difficulty = 7.5

            stability = 1.0
            if c_mnem:
                if isinstance(c_mnem, dict) and c_mnem.get("keyword"):
                    stability = 1.5
                elif isinstance(c_mnem, str) and c_mnem.strip():
                    stability = 1.5

            c_type = (c.get("content_type") if isinstance(c, dict) else getattr(c, "content_type", None)) or ("cloze" if "{{c" in c_text_clean else "text")
            new_card = Card(
                phrase_id=target_phrase.id,
                user_id=user_id,
                subject=clean_sub,
                text=c_text_clean,
                secondary_text=c_sec,
                translation=c_trans,
                example=c_ex,
                difficulty=difficulty,
                stability=stability,
                state=0,
                mnemonic=c_mnem,
                content_type=c_type,
                next_review=now
            )
            db.add(new_card)
            existing_map[c_key] = new_card  # предотвращаем дубли внутри пачки
            cards_created += 1

    return cards_created, cards_updated, clean_sub, clean_title

# --- 1. ВЫДАЧА АРХИВА КАРТОЧЕК ТЕКУЩЕГО ПОЛЬЗОВАТЕЛЯ ---
@router.get("/data/cards")
async def get_all_cards(
    subject: str = Query("all"), 
    page: int = Query(1, ge=1), 
    limit: int = Query(1000, ge=1, le=10000), 
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    stmt = select(Card).filter(Card.user_id == current_user)
    if subject != 'all': 
        stmt = stmt.filter(Card.subject == subject)
    
    count_stmt = select(func.count()).select_from(stmt.subquery())
    count_res = await db.execute(count_stmt)
    total = count_res.scalar_one()
    
    offset = (page - 1) * limit
    stmt = stmt.offset(offset).limit(limit)
    cards_res = await db.execute(stmt)
    cards = cards_res.scalars().all()
    
    return {
        "total": total, 
        "page": page, 
        "limit": limit,
        "cards": [
            {
                "id": c.id, 
                "text": c.text, 
                "secondary_text": c.secondary_text if c.secondary_text else "", 
                "translation": c.translation, 
                "state": c.state, 
                "subject": c.subject, 
                "example": c.example if c.example else "",
                "mnemonic": c.mnemonic
            } 
            for c in cards
        ]
    }

@router.get("/data/cards/export")
async def export_cards_json(
    subject: str = Query("all"),
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    """
    Экспорт карточек текущего пользователя в эталонном формате пресета .json.
    Идеально подходит для скачивания колоды и последующей раздачи всем участникам.
    """
    stmt = select(Card).filter(Card.user_id == current_user)
    if subject != "all":
        stmt = stmt.filter(Card.subject == subject)
    stmt = stmt.order_by(Card.id.asc())

    res = await db.execute(stmt)
    cards = res.scalars().all()

    # Если у пользователя нет карт, но он админ/dev в браузере — проверяем default_user
    if not cards and current_user != "default_user":
        stmt_def = select(Card).filter(Card.user_id == "default_user")
        if subject != "all":
            stmt_def = stmt_def.filter(Card.subject == subject)
        stmt_def = stmt_def.order_by(Card.id.asc())
        res_def = await db.execute(stmt_def)
        cards = res_def.scalars().all()

    phrase_title = "Судоустройство: Основной курс" if subject == "sudoustroystvo" else f"Курс: {subject}"
    if cards:
        p_stmt = select(Phrase.text).filter(Phrase.user_id == cards[0].user_id)
        if subject != "all":
            p_stmt = p_stmt.filter(Phrase.subject == subject)
        found_title = (await db.execute(p_stmt)).scalar()
        if found_title:
            phrase_title = found_title

    payload = {
        "phrase_title": phrase_title,
        "subject_slug": subject if subject != "all" else (cards[0].subject if cards else "sudoustroystvo"),
        "total_cards": len(cards),
        "cards": [
            {
                "text": c.text,
                "secondary_text": c.secondary_text or "",
                "translation": c.translation,
                "example": c.example or "",
                "mnemonic": c.mnemonic
            }
            for c in cards
        ]
    }

    json_bytes = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
    sub_tag = subject if subject != "all" else "all_subjects"
    filename = f"grinder_deck_{sub_tag}.json"

    return Response(
        content=json_bytes,
        media_type="application/json; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-cache"
        }
    )

@router.post("/data/cards/share")
async def share_cards_deck(
    payload: ShareDeckIn,
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    """
    Генерирует уникальный ключ/диплинк для вирусного шеринга колоды карточек с друзьями в Telegram.
    """
    sub = payload.subject.strip().lower()
    if not sub or sub == "all":
        raise HTTPException(status_code=400, detail="Укажите конкретный предмет для шеринга.")

    stmt = select(Card).filter(Card.user_id == current_user, Card.subject == sub).order_by(Card.id.asc())
    cards = (await db.execute(stmt)).scalars().all()
    if not cards and current_user != "default_user":
        stmt_def = select(Card).filter(Card.user_id == "default_user", Card.subject == sub).order_by(Card.id.asc())
        cards = (await db.execute(stmt_def)).scalars().all()

    if not cards:
        raise HTTPException(status_code=404, detail="В этой колоде пока нет карточек для шеринга.")

    p_stmt = select(Phrase.text).filter(Phrase.user_id == cards[0].user_id, Phrase.subject == sub)
    found_title = (await db.execute(p_stmt)).scalar() or f"Колода: {sub}"

    share_key = secrets.token_hex(4)
    deck_data = {
        "phrase_title": found_title,
        "subject_slug": sub,
        "created_by": current_user,
        "created_at": datetime.utcnow().isoformat(),
        "total_cards": len(cards),
        "cards": [
            {
                "text": c.text,
                "secondary_text": c.secondary_text or "",
                "translation": c.translation,
                "example": c.example or "",
                "mnemonic": c.mnemonic,
                "content_type": getattr(c, "content_type", "text")
            }
            for c in cards
        ]
    }

    shares_dir = Path("app/static/presets/shares")
    shares_dir.mkdir(parents=True, exist_ok=True)
    share_file = shares_dir / f"{share_key}.json"
    share_file.write_text(json.dumps(deck_data, ensure_ascii=False, indent=2), encoding="utf-8")

    bot_username = getattr(settings, "TELEGRAM_BOT_USERNAME", "") or "DataGrinderBot"
    bot_clean = str(bot_username).lstrip("@")
    tg_link = f"https://t.me/{bot_clean}?start=deck_{share_key}"

    return {
        "status": "success",
        "share_key": share_key,
        "subject": sub,
        "theme": found_title,
        "title": found_title,
        "cards_count": len(cards),
        "total_cards": len(cards),
        "tg_link": tg_link,
        "share_url": tg_link
    }

# --- 2. АНАЛИТИКА И ДАШБОРД ТЕКУЩЕГО ПОЛЬЗОВАТЕЛЯ ---
@router.get("/stats/dashboard")
async def get_analytics(
    subject: str = Query("all"), 
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    card_stmt = select(Card).filter(Card.user_id == current_user)
    if subject != 'all': 
        card_stmt = card_stmt.filter(Card.subject == subject)
    card_res = await db.execute(card_stmt)
    cards = card_res.scalars().all()
    
    states_dict = {0: 0, 1: 0, 2: 0, 3: 0}
    for c in cards: 
        states_dict[c.state if c.state in states_dict else 0] += 1
        
    total_cards = len(cards)
    progress_percent = round((states_dict[2] / total_cards) * 100) if total_cards > 0 else 0

    # Расчет Retention Rate за 30 дней для конкретного пользователя
    one_month_ago = datetime.utcnow() - timedelta(days=30)
    
    if subject != 'all':
        log_stmt = select(ReviewLog).join(Card, ReviewLog.card_id == Card.id).filter(
            ReviewLog.user_id == current_user,
            ReviewLog.review_time >= one_month_ago, 
            Card.subject == subject
        )
    else:
        log_stmt = select(ReviewLog).filter(
            ReviewLog.user_id == current_user,
            ReviewLog.review_time >= one_month_ago
        )
        
    log_res = await db.execute(log_stmt)
    logs = log_res.scalars().all()
    total_reviews = len(logs)
    successful_reviews = sum(1 for r in logs if r.rating > 1)
    retention_rate = round((successful_reviews / total_reviews) * 100, 1) if total_reviews > 0 else 0.0

    # Расчет ударного режима (Streak) для конкретного пользователя
    streak_stmt = select(func.date(ReviewLog.review_time)).filter(
        ReviewLog.user_id == current_user
    ).distinct().order_by(func.date(ReviewLog.review_time).desc()).limit(30)
    
    streak_res = await db.execute(streak_stmt)
    active_days = streak_res.all()
    
    streak = 0
    dates_set = {datetime.strptime(str(d[0]), "%Y-%m-%d").date() if isinstance(d[0], str) else d[0] for d in active_days}
    current_date = datetime.utcnow().date()
    if current_date not in dates_set: 
        current_date -= timedelta(days=1)
    while current_date in dates_set:
        streak += 1
        current_date -= timedelta(days=1)

    # Тематическая раскладка матрицы знаний текущего пользователя
    breakdown = []
    if subject == "all":
        by_sub = defaultdict(list)
        for c in cards:
            if c.subject:
                by_sub[c.subject].append(c)
        for sub, sub_cards in by_sub.items():
            sub_total = len(sub_cards)
            sub_review = sum(1 for c in sub_cards if c.state == 2)
            breakdown.append({
                "label": sub.upper(), 
                "progress": round((sub_review / sub_total) * 100) if sub_total > 0 else 0
            })
    else:
        phrase_stmt = select(Phrase).filter(Phrase.subject == subject, Phrase.user_id == current_user)
        phrase_res = await db.execute(phrase_stmt)
        phrases = phrase_res.scalars().all()
        for phrase in phrases:
            phrase_cards = [c for c in cards if c.phrase_id == phrase.id]
            p_total = len(phrase_cards)
            p_review = sum(1 for c in phrase_cards if c.state == 2)
            breakdown.append({
                "label": phrase.text, 
                "progress": round((p_review / p_total) * 100) if p_total > 0 else 0
            })

    # Проверяем, пройден ли опрос сегодня именно этим пользователем
    msk_now = datetime.utcnow() + timedelta(hours=3)
    msk_today_start = msk_now.replace(hour=0, minute=0, second=0, microsecond=0)
    utc_today_start = msk_today_start - timedelta(hours=3)
    
    survey_stmt = select(DailySession).filter(
        DailySession.user_id == current_user,
        DailySession.timestamp >= utc_today_start
    )
    survey_res = await db.execute(survey_stmt)
    survey_completed = survey_res.scalars().first() is not None

    # Карты к вечеру для текущего пользователя
    if msk_now.hour < 21:
        evening_msk = msk_now.replace(hour=21, minute=0, second=0, microsecond=0)
    else:
        evening_msk = (msk_now + timedelta(days=1)).replace(hour=21, minute=0, second=0, microsecond=0)
    evening_utc = evening_msk - timedelta(hours=3)
    
    evening_stmt = select(func.count(Card.id)).filter(
        Card.user_id == current_user,
        Card.state.in_([1, 2, 3]),
        Card.next_review <= evening_utc
    )
    if subject != 'all':
        evening_stmt = evening_stmt.filter(Card.subject == subject)
    evening_res = await db.execute(evening_stmt)
    due_evening = evening_res.scalar() or 0

    # Проверяем карточки, требующие повторения прямо сейчас (REV просроченные + LRN краткосрочные)
    now_utc = datetime.utcnow()
    due_stmt = select(func.count(Card.id)).filter(
        Card.user_id == current_user,
        (
            (Card.state == 2) & (Card.next_review <= now_utc)
        ) | (
            Card.state.in_([1, 3])
        )
    )
    if subject != 'all':
        due_stmt = due_stmt.filter(Card.subject == subject)
    due_res = await db.execute(due_stmt)
    due_reviews_now = due_res.scalar() or 0

    # Проверяем участие в эксперименте и рассчитываем дневную квоту новых карт
    session_stmt = select(UserSession).filter(UserSession.user_id == current_user)
    sess_res = await db.execute(session_stmt)
    user_sess = sess_res.scalars().first()
    is_participant = bool(user_sess and user_sess.is_experiment_participant)
    phase = user_sess.experiment_phase if user_sess else 1

    if is_participant and phase == 1:
        daily_new_limit = settings.EXPERIMENT_DAILY_LIMIT
    else:
        setting_res = await db.execute(select(UserSetting).filter(UserSetting.user_id == current_user))
        user_setting = setting_res.scalar_one_or_none()
        user_daily_limit = user_setting.daily_limit if user_setting else 10
        subject_limits = user_setting.subject_limits if (user_setting and user_setting.subject_limits) else {}
        daily_new_limit = subject_limits.get(subject, user_daily_limit)

    # Сколько новых карточек изучено сегодня
    today_start = datetime.utcnow().replace(hour=0, minute=0, second=0, microsecond=0)
    new_today_stmt = select(ReviewLog.id).join(Card, ReviewLog.card_id == Card.id).filter(
        ReviewLog.user_id == current_user,
        ReviewLog.state == 0,
        ReviewLog.review_time >= today_start
    )
    if subject != 'all':
        new_today_stmt = new_today_stmt.filter(Card.subject == subject)
    new_today_res = await db.execute(new_today_stmt)
    already_learned_today = len(new_today_res.scalars().all())

    unlearned_in_deck = states_dict[0]
    allowed_new_count = max(0, daily_new_limit - already_learned_today)
    new_remaining_today = min(allowed_new_count, unlearned_in_deck)

    # 1. Распределение зрелости колоды по FSRS стабильности:
    # Хрупкие (S < 7), Развивающиеся (7 <= S < 30), Зрелые (30 <= S < 180), Долговременные (S >= 180)
    maturity = {
        "new": states_dict[0],
        "fragile": 0,
        "developing": 0,
        "mature": 0,
        "mastered": 0
    }
    for c in cards:
        if c.state != 0:
            s_val = c.stability or 0.0
            if s_val < 7.0:
                maturity["fragile"] += 1
            elif s_val < 30.0:
                maturity["developing"] += 1
            elif s_val < 180.0:
                maturity["mature"] += 1
            else:
                maturity["mastered"] += 1

    # 2. Тепловая карта активности (Heatmap) за последние 60 дней
    sixty_days_ago = datetime.utcnow() - timedelta(days=60)
    heatmap_stmt = select(func.date(ReviewLog.review_time), func.count(ReviewLog.id)).filter(
        ReviewLog.user_id == current_user,
        ReviewLog.review_time >= sixty_days_ago
    ).group_by(func.date(ReviewLog.review_time))
    heatmap_res = await db.execute(heatmap_stmt)
    heatmap_data = {str(row[0]): row[1] for row in heatmap_res.all() if row[0]}

    return {
        "cards_new": states_dict[0], 
        "cards_learning": states_dict[1] + states_dict[3], 
        "cards_review": states_dict[2],
        "total_cards": total_cards,
        "due_reviews_now": due_reviews_now,
        "new_remaining_today": new_remaining_today,
        "daily_new_limit": daily_new_limit,
        "already_learned_today": already_learned_today,
        "unlearned_in_deck": unlearned_in_deck,
        "progress_percent": f"{progress_percent}%", 
        "retention_rate_30d": f"{retention_rate}%", 
        "streak_days": streak, 
        "breakdown": breakdown,
        "survey_completed": survey_completed,
        "due_evening": due_evening,
        "maturity": maturity,
        "heatmap": heatmap_data
    }

# --- 3. НАСТРОЙКИ ПОЛЬЗОВАТЕЛЯ ИЗ ТАБЛИЦЫ БД ---
@router.get("/config")
async def get_config(
    subject: str = Query("all"),
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    setting_res = await db.execute(select(UserSetting).filter(UserSetting.user_id == current_user))
    setting = setting_res.scalar_one_or_none()
    
    daily_limit = 10
    subject_limits = {}
    assoc_pref = "acoustic"
    target_retention = 0.9
    
    if setting:
        daily_limit = setting.daily_limit
        subject_limits = setting.subject_limits or {}
        assoc_pref = setting.assoc_preference or "acoustic"
        target_retention = setting.target_retention or 0.9
        
    current_subject_limit = subject_limits.get(subject, daily_limit)
    return {
        "daily_limit": current_subject_limit, 
        "focus_mode_default": False,
        "assoc_preference": assoc_pref,
        "target_retention": target_retention
    }

@router.post("/config")
async def update_config(
    payload: ConfigUpdate, 
    subject: str = Query("all"),
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    setting_res = await db.execute(select(UserSetting).filter(UserSetting.user_id == current_user))
    setting = setting_res.scalar_one_or_none()
    
    if not setting:
        setting = UserSetting(
            user_id=current_user,
            daily_limit=payload.daily_limit,
            subject_limits={subject: payload.daily_limit}
        )
        db.add(setting)
    else:
        subject_limits = dict(setting.subject_limits or {})
        subject_limits[subject] = payload.daily_limit
        setting.subject_limits = subject_limits
        if subject == "all":
            setting.daily_limit = payload.daily_limit
        if payload.target_retention is not None:
            setting.target_retention = payload.target_retention
        if payload.assoc_preference is not None:
            setting.assoc_preference = payload.assoc_preference

    await db.commit()
    return {"status": "updated", "config": {"daily_limit": payload.daily_limit}}

# --- 4. ИИ-КОНВЕЙЕР ИМПОРТА И ПЕСОЧНИЦА (STAGING SANDBOX) ---
@router.post("/config/import")
async def import_raw_text(
    payload: ImportIn, 
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    if not payload.text.strip(): 
        return {"status": "error", "message": "Входящий текст пуст."}
    
    target_sub = payload.subject.strip().lower()
    if not target_sub:
        return {"status": "error", "message": "Целевой предмет не выбран. Выберите предмет из списка или укажите новый."}

    # Если выбрана отложенная обработка «Ночной Грайнд» (-50% стоимости)
    # ИЛИ объем текста превышает 35 000 знаков (автоматический фоновый режим во избежание таймаута)
    is_too_large = len(payload.text.strip()) > 35000
    if payload.is_deferred or is_too_large:
        is_offpeak = is_deepseek_offpeak()
        if is_offpeak:
            # Скидка 50% УЖЕ действует прямо сейчас (19:30-03:30 МСК)!
            # Нарезка начинается немедленно, задача не задерживается
            effective_deferred = False
            is_immediate = True
            if is_too_large:
                msg = f"🔥 Скидка 50% активна прямо сейчас! Объемный текст ({len(payload.text.strip())} знаков) взят в фоновую нарезку со скидкой 50%."
            else:
                msg = f"🔥 Скидка 50% активна прямо сейчас! Материал передан в немедленную фоновую обработку."
        else:
            if payload.is_deferred:
                # В дневное время пользователь явно выбрал скидку 50% в 19:30 МСК
                effective_deferred = True
                is_immediate = False
                msg = "Материал принят в очередь «Ночной Грайнд». Обработка начнется в 19:30 по МСК (со скидкой 50%). Мы уведомим вас о готовности!"
            else:
                # Дневное время, но пользователь выбрал генерацию сейчас: фоновый запуск без откладывания
                effective_deferred = False
                is_immediate = True
                msg = f"Объемный материал ({len(payload.text.strip())} знаков) взят в немедленную фоновую обработку по дневному тарифу."

        # Извлекаем осмысленное имя темы (пропуская служебные технические разделители OCR)
        meaningful_lines = [
            l.strip() for l in payload.text.strip().split("\n")
            if l.strip() and not l.strip().startswith("===") and not l.strip().startswith("---")
        ]
        if meaningful_lines:
            extracted_theme = meaningful_lines[0][:40].strip()
        else:
            first_line = payload.text.strip().split("\n")[0][:40].strip()
            extracted_theme = first_line.replace("===", "").strip() or "Новый блок знаний"

        job = GenerationJob(
            user_id=current_user,
            telegram_id=current_user if current_user.isdigit() else None,
            subject=target_sub,
            theme=extracted_theme,
            raw_text=payload.text.strip(),
            granularity_mode=payload.granularity_mode,
            density=payload.density,
            volume=payload.volume,
            custom_instruction=payload.custom_instruction.strip(),
            is_deferred=effective_deferred,
            status="pending"
        )
        db.add(job)
        await db.commit()
        await db.refresh(job)

        return {
            "status": "queued",
            "job_id": job.id,
            "subject": target_sub,
            "theme": job.theme,
            "is_immediate": is_immediate,
            "is_offpeak": is_offpeak,
            "message": msg
        }


    try: 
        parsed_data = await parse_raw_text(
            payload.text,
            target_subject=target_sub,
            density=payload.density,
            volume=payload.volume,
            priority=payload.priority,
            preference=payload.assoc_preference,
            granularity_mode=payload.granularity_mode,
            custom_instruction=payload.custom_instruction
        )
    except Exception as e: 
        import traceback
        traceback.print_exc()
        err_msg = str(e)
        print(f"[ERROR /api/config/import] {err_msg}")
        if "RESOURCE_EXHAUSTED" in err_msg or "429" in err_msg:
            return {"status": "error", "message": f"Лимит или баланс ИИ-провайдера ({settings.AI_PROVIDER}) исчерпан (429). Проверьте баланс или ключ API."}
        return {"status": "error", "message": f"Ошибка ИИ-генератора ({settings.AI_PROVIDER}): {err_msg}"}
        
    if "error" in parsed_data: 
        err_msg = str(parsed_data["error"])
        print(f"[ERROR /api/config/import from parsed_data] {err_msg}")
        if "RESOURCE_EXHAUSTED" in err_msg or "429" in err_msg:
            return {"status": "error", "message": f"Лимит или баланс ИИ-провайдера ({settings.AI_PROVIDER}) исчерпан (429). Проверьте баланс или ключ API."}
        return {"status": "error", "message": err_msg}
        
    subject_slug = target_sub
    phrase_title = parsed_data.get("phrase_title", "Новый блок знаний")
    cards = parsed_data.get("cards", [])

    # Если включен режим Staging (по умолчанию True для фронтенда), возвращаем карточки в Песочницу!
    if not payload.commit_now:
        return {
            "status": "staging",
            "subject": subject_slug,
            "theme": phrase_title,
            "cards": cards
        }

    # Прямое сохранение, если запрошено
    try:
        cards_created, clean_sub, clean_title = await save_cards_to_database(
            cards_data=cards, 
            subject_slug=subject_slug, 
            phrase_title=phrase_title, 
            user_id=current_user, 
            db=db
        )
        if cards_created > 0:
            await db.commit()
            return {"status": "success", "subject": clean_sub, "theme": clean_title, "cards_count": cards_created}
        else:
            await db.rollback()
            return {"status": "error", "message": "ИИ не смог нарезать карточки."}
    except Exception as e:
        await db.rollback()
        raise HTTPException(status_code=500, detail=f"Ошибка БД: {str(e)}")

# --- 4.1 ФИКСАЦИЯ ОДОБРЕННЫХ КАРТОЧЕК ИЗ ПЕСОЧНИЦЫ (STAGING COMMIT) ---
@router.post("/config/import/commit")
async def commit_staging_cards(
    payload: StagingCommitIn,
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    if not payload.cards:
        raise HTTPException(status_code=400, detail="Список одобренных карточек пуст.")

    try:
        cards_created, clean_sub, clean_title = await save_cards_to_database(
            cards_data=[c.dict() for c in payload.cards],
            subject_slug=payload.subject,
            phrase_title=payload.theme,
            user_id=current_user,
            db=db
        )
        # Если карточки были импортированы из фоновой задачи, помечаем её завершенной
        if payload.job_id:
            stmt_job = select(GenerationJob).filter(GenerationJob.id == payload.job_id, GenerationJob.user_id == current_user)
            res_job = await db.execute(stmt_job)
            job = res_job.scalar_one_or_none()
            if job:
                job.status = "completed"
                job.cards_count = cards_created
                job.result_cards_json = None  # Освобождаем место в БД после фиксации в карточки
        await db.commit()
        return {
            "status": "success", 
            "subject": clean_sub, 
            "theme": clean_title, 
            "cards_count": cards_created
        }
    except Exception as e:
        await db.rollback()
        raise HTTPException(status_code=500, detail=f"Ошибка сохранения из песочницы: {str(e)}")

# --- 4.2 ЗАГРУЗКА ПАЧЕК ФАЙЛОВ НА КОДОВОМ УРОВНЕ (PDF, TXT, MD, CSV) ---
@router.post("/config/import/file")
async def import_file_at_code_level(
    request: Request,
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    form = await request.form()
    
    # Извлекаем все переданные файлы (поддерживаем ключи 'files' и 'file', единичные и множественные)
    raw_files = form.getlist("files")
    if not raw_files and "file" in form:
        raw_files = [form.get("file")]
    
    upload_list: list[UploadFile] = [
        f for f in raw_files 
        if hasattr(f, "filename") and f.filename
    ]
    if not upload_list:
        raise HTTPException(status_code=400, detail="Не передано ни одного файла.")

    target_sub = str(form.get("subject", "")).strip().lower()
    if not target_sub:
        raise HTTPException(status_code=400, detail="Целевой предмет не выбран. Выберите предмет из списка или укажите новый.")

    density = str(form.get("density", "medium"))
    volume = str(form.get("volume", "auto"))
    priority = str(form.get("priority", "balanced"))
    assoc_preference = str(form.get("assoc_preference", "acoustic"))
    granularity_mode = str(form.get("granularity_mode", "atomic"))
    custom_instruction = str(form.get("custom_instruction", "")).strip()
    commit_now = str(form.get("commit_now", "")).lower() == "true"
    is_deferred = str(form.get("is_deferred", "")).lower() == "true"

    all_cards: list[dict] = []
    all_extracted_texts: list[str] = []
    file_titles: list[str] = []

    for up_file in upload_list:
        filename = (up_file.filename or "file").lower()
        file_titles.append(up_file.filename or "файл")
        contents = await up_file.read()
        extracted_text = ""

        # 1. Формат PDF: извлечение через pypdf (с безопасным ограничением на объем)
        if filename.endswith(".pdf"):
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(contents))
                pages_text = []
                total_pages = len(reader.pages)
                # Безопасно извлекаем до 500 страниц книги
                max_pages = min(total_pages, 500)
                for idx in range(max_pages):
                    txt = reader.pages[idx].extract_text() or ""
                    if txt.strip():
                        pages_text.append(f"--- {up_file.filename}: Стр. {idx+1} ---\n{txt.strip()}")
                
                extracted_text = "\n\n".join(pages_text)
                print(f"[PDF Import] {up_file.filename}: извлечено {len(pages_text)} из {total_pages} страниц ({len(extracted_text)} знаков).")

            except Exception as e:
                print(f"[WARN] Ошибка чтения PDF {up_file.filename}: {e}")

        # 2. Формат CSV / TSV: прямой парсинг
        elif filename.endswith(".csv") or filename.endswith(".tsv"):
            delimiter = "\t" if filename.endswith(".tsv") else ","
            try:
                text_stream = io.StringIO(contents.decode("utf-8-sig", errors="ignore"))
                reader = csv.reader(text_stream, delimiter=delimiter)
                for row in reader:
                    if len(row) >= 2 and row[0].strip() and row[1].strip():
                        all_cards.append({
                            "text": row[0].strip(),
                            "secondary_text": row[2].strip() if len(row) > 2 else "",
                            "translation": row[1].strip(),
                            "example": "",
                            "initial_difficulty_tier": "medium",
                            "mnemonic": None
                        })
            except Exception as e:
                print(f"[WARN] Ошибка чтения CSV {up_file.filename}: {e}")

        # 3. Обычный текстовый или Markdown файл
        elif filename.endswith(".txt") or filename.endswith(".md"):
            try:
                extracted_text = contents.decode("utf-8-sig", errors="ignore")
            except Exception as e:
                print(f"[WARN] Ошибка чтения TXT {up_file.filename}: {e}")

        # 4. Формат DOCX (Microsoft Word / Google Docs)
        elif filename.endswith(".docx"):
            try:
                import docx
                doc = docx.Document(io.BytesIO(contents))
                doc_paragraphs = []
                for p in doc.paragraphs:
                    p_text = p.text.strip()
                    if p_text:
                        doc_paragraphs.append(p_text)
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if row_cells:
                            doc_paragraphs.append(" | ".join(row_cells))
                extracted_text = "\n\n".join(doc_paragraphs)
                print(f"[DOCX Import] {up_file.filename}: извлечено {len(doc_paragraphs)} параграфов ({len(extracted_text)} знаков).")
            except Exception as e:
                print(f"[WARN] Ошибка чтения DOCX {up_file.filename}: {e}")

        # 5. Формат PPTX (Microsoft PowerPoint / Слайды лекций)
        elif filename.endswith(".pptx"):
            try:
                from pptx import Presentation
                prs = Presentation(io.BytesIO(contents))
                slides_text = []
                for s_idx, slide in enumerate(prs.slides, 1):
                    slide_lines = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            slide_lines.append(shape.text.strip())
                    if slide_lines:
                        slides_text.append(f"--- {up_file.filename}: Слайд {s_idx} ---\n" + "\n".join(slide_lines))
                extracted_text = "\n\n".join(slides_text)
                print(f"[PPTX Import] {up_file.filename}: извлечено {len(slides_text)} слайдов ({len(extracted_text)} знаков).")
            except Exception as e:
                print(f"[WARN] Ошибка чтения PPTX {up_file.filename}: {e}")

        if extracted_text.strip():
            all_extracted_texts.append(f"=== ДОКУМЕНТ: {up_file.filename} ===\n" + extracted_text.strip())

    # Если включен режим «Ночной Грайнд» (-50% стоимости) для документов
    # ИЛИ объем документов превышает 30 000 знаков / более 1 документа (автоматический фоновый режим)
    if all_extracted_texts:
        combined_text = "\n\n".join(all_extracted_texts)
        is_large = len(combined_text) > 30000 or len(upload_list) > 1

        if is_deferred or is_large:
            is_offpeak = is_deepseek_offpeak()
            if is_offpeak:
                # В часы скидок (19:30-03:30 МСК) скидка 50% УЖЕ действует прямо сейчас!
                # Задачи НЕ откладываются на потом, а запускаются фоновым воркером немедленно со скидкой 50%
                effective_deferred = False
                is_immediate = True
                if is_large:
                    msg = f"🔥 Скидка 50% активна прямо сейчас! Крупный документ ({len(combined_text)} знаков) взят в фоновую нарезку со скидкой 50%."
                else:
                    msg = f"🔥 Скидка 50% активна! Файлы ({len(upload_list)} шт.) переданы в немедленную фоновую обработку."
            else:
                # В дневное время:
                if is_deferred:
                    # Пользователь явно выбрал отложить до 19:30 для скидки 50%
                    effective_deferred = True
                    is_immediate = False
                    msg = f"Файлы ({len(upload_list)} шт.) поставлены в очередь «Ночной Грайнд». Обработка начнется в 19:30 по МСК (со скидкой 50%)."
                else:
                    # Пользователь нажал "Создать сейчас" (дневной тариф):
                    # Крупные материалы отправляются в фоновый воркер во избежание HTTP таймаута, но запускаются НЕМЕДЛЕННО
                    effective_deferred = False
                    is_immediate = True
                    msg = f"Крупный документ ({len(combined_text)} знаков) взят в немедленную фоновую нарезку по дневному тарифу."

            theme_name = f"Пакетный импорт ({len(upload_list)} док.): {', '.join(file_titles[:2])}"
            if len(file_titles) > 2:
                theme_name += f" и ещё {len(file_titles) - 2}"

            job = GenerationJob(
                user_id=current_user,
                telegram_id=current_user if current_user.isdigit() else None,
                subject=target_sub,
                theme=theme_name,
                raw_text=combined_text,
                granularity_mode=granularity_mode,
                density=density,
                volume=volume,
                custom_instruction=custom_instruction,
                is_deferred=effective_deferred,
                status="pending"
            )
            db.add(job)
            await db.commit()
            await db.refresh(job)

            return {
                "status": "queued",
                "job_id": job.id,
                "subject": target_sub,
                "theme": theme_name,
                "is_immediate": is_immediate,
                "is_offpeak": is_offpeak,
                "message": msg
            }


    # Если были текстовые/PDF документы, прогоняем через ИИ мгновенно
    if all_extracted_texts:
        combined_text = "\n\n".join(all_extracted_texts)
        try:
            parsed_data = await parse_raw_text(
                combined_text,
                target_subject=target_sub,
                density=density,
                volume=volume,
                priority=priority,
                preference=assoc_preference,
                granularity_mode=granularity_mode,
                custom_instruction=custom_instruction
            )
            if "cards" in parsed_data and isinstance(parsed_data["cards"], list):
                all_cards.extend(parsed_data["cards"])
        except Exception as e:
            import traceback
            traceback.print_exc()
            print(f"[ERROR /api/config/import/file] {e}")
            if not all_cards:
                raise HTTPException(status_code=500, detail=f"Ошибка ИИ при структурировании документов ({settings.AI_PROVIDER}): {str(e)}")

    if not all_cards:
        raise HTTPException(status_code=400, detail="Не удалось извлечь карточки из переданных файлов.")

    subject_slug = target_sub
    theme_name = f"Пакетный импорт ({len(upload_list)} док.): {', '.join(file_titles[:2])}"
    if len(file_titles) > 2:
        theme_name += f" и ещё {len(file_titles) - 2}"

    if not commit_now:
        return {
            "status": "staging",
            "subject": subject_slug,
            "theme": theme_name,
            "cards": all_cards
        }
    else:
        cards_created, clean_sub, clean_title = await save_cards_to_database(
            cards_data=all_cards,
            subject_slug=subject_slug,
            phrase_title=theme_name,
            user_id=current_user,
            db=db
        )
        await db.commit()
        return {"status": "success", "subject": clean_sub, "theme": clean_title, "cards_count": cards_created}

# --- 4.2.0 ВЫГРУЗКА КАРТОЧЕК ИЗ ЗАДАЧИ В ПЕСОЧНИЦУ (STAGING SANDBOX) ---
@router.get("/config/import/staging/job/{job_id}")
async def get_staging_job_cards(
    job_id: int,
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    """Возвращает сформированные карточки фоновой задачи для разбора в Песочнице."""
    stmt = select(GenerationJob).filter(GenerationJob.id == job_id, GenerationJob.user_id == current_user)
    res = await db.execute(stmt)
    job = res.scalar_one_or_none()
    if not job:
        raise HTTPException(status_code=404, detail="Задача не найдена или нет прав доступа.")

    if not job.result_cards_json:
        if job.status in ("pending", "processing"):
            raise HTTPException(status_code=400, detail="Карточки ещё нарезаются ИИ в фоновом режиме. Пожалуйста, подождите завершения.")
        elif job.status == "failed":
            raise HTTPException(status_code=400, detail=f"Ошибка обработки: {job.error_message or 'Неизвестная ошибка'}")
        elif job.status == "completed":
            raise HTTPException(status_code=400, detail="Карточки из этой задачи уже были сохранены в базу знаний.")
        else:
            raise HTTPException(status_code=400, detail="Для этой задачи нет готовых карточек.")

    try:
        cards = json.loads(job.result_cards_json)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ошибка распаковки карточек: {str(e)}")

    return {
        "status": "staging",
        "job_id": job.id,
        "subject": job.subject,
        "theme": job.theme,
        "cards": cards
    }

# --- 4.2.1 ОЧЕРЕДЬ НОЧНОГО ГРАЙНДА И ФОНОВЫХ ЗАДАЧ ---
@router.get("/config/import/queue")
async def get_import_queue(
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    """Возвращает статус задач пользователя в очереди Ночного Грайндера и фоновых задач."""
    stmt = (
        select(GenerationJob)
        .filter(GenerationJob.user_id == current_user)
        .order_by(GenerationJob.id.desc())
        .limit(10)
    )
    res = await db.execute(stmt)
    jobs = res.scalars().all()
    return {
        "status": "success",
        "jobs": [
            {
                "id": j.id,
                "subject": j.subject,
                "theme": j.theme,
                "status": j.status,
                "cards_count": j.cards_count,
                "has_cards": bool(j.result_cards_json),
                "error_message": j.error_message,
                "created_at": j.created_at.isoformat() if j.created_at else None,
                "processed_at": j.processed_at.isoformat() if j.processed_at else None
            }
            for j in jobs
        ]
    }

@router.delete("/config/import/queue/{job_id}")
async def cancel_import_queue_job(
    job_id: int,
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    """Отменяет задачу пользователя в очереди ночной генерации."""
    stmt = select(GenerationJob).filter(GenerationJob.id == job_id, GenerationJob.user_id == current_user)
    res = await db.execute(stmt)
    job = res.scalar_one_or_none()
    if not job:
        raise HTTPException(status_code=404, detail="Задача не найдена или нет прав доступа.")
    
    if job.status in ("completed", "failed", "cancelled"):
        raise HTTPException(status_code=400, detail=f"Нельзя отменить задачу в статусе {job.status}.")
        
    job.status = "cancelled"
    job.error_message = "Отменено пользователем"
    await db.commit()
    return {"status": "success", "message": f"Задача #{job_id} успешно отменена."}

# --- 4.3 ИМПОРТ ГОТОВОЙ БИБЛИОТЕКИ ---
@router.post("/config/import/preset")
async def import_preset_library(
    payload: PresetImportIn, 
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    import os
    import json
    
    preset_name = payload.preset_name.strip().lower()
    if not re.match(r'^[a-z0-9_]+$', preset_name):
        raise HTTPException(status_code=400, detail="Недопустимое имя пресета.")
    preset_filename = f"{preset_name}.json"
    preset_path = os.path.join("app", "static", "presets", preset_filename)
    
    if not os.path.exists(preset_path):
        raise HTTPException(status_code=404, detail="Библиотека не найдена.")
        
    try:
        with open(preset_path, "r", encoding="utf-8") as f:
            parsed_data = json.load(f)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ошибка чтения файла пресета: {str(e)}")
        
    subject_slug = parsed_data.get("subject_slug", "generic").lower()
    phrase_title = parsed_data.get("phrase_title", "Новый блок знаний")
    cards = parsed_data.get("cards", [])

    if not payload.commit_now:
        return {
            "status": "staging",
            "subject": subject_slug,
            "theme": phrase_title,
            "cards": cards
        }
        
    try:
        cards_created, clean_sub, clean_title = await save_cards_to_database(
            cards_data=cards,
            subject_slug=subject_slug,
            phrase_title=phrase_title,
            user_id=current_user,
            db=db
        )
        if cards_created > 0:
            await db.commit()
            return {"status": "success", "subject": clean_sub, "theme": clean_title, "cards_count": cards_created}
        else:
            await db.rollback()
            return {"status": "error", "message": "В библиотеке нет валидных карточек."}
    except Exception as e:
        await db.rollback()
        raise HTTPException(status_code=500, detail=f"Ошибка БД: {str(e)}")

# --- 5. РУЧНОЕ СОЗДАНИЕ И РЕДАКТИРОВАНИЕ КАРТОЧЕК ---
@router.post("/management/cards")
async def create_manual_card(
    payload: ManualCardIn,
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    if not payload.text.strip() or not payload.translation.strip():
        raise HTTPException(status_code=400, detail="Лицевая сторона и перевод обязательны.")

    clean_sub = payload.subject.strip().lower() or "generic"
    clean_title = payload.phrase_title.strip() or "Пользовательские карточки"
    
    phrase_res = await db.execute(
        select(Phrase).filter(Phrase.text == clean_title, Phrase.subject == clean_sub, Phrase.user_id == current_user)
    )
    phrase = phrase_res.scalar_one_or_none()
    if not phrase:
        phrase = Phrase(text=clean_title, subject=clean_sub, user_id=current_user)
        db.add(phrase)
        await db.flush()

    mnemonic_json = None
    if payload.mnemonic_keyword or payload.mnemonic_cue:
        mnemonic_json = {
            "keyword": payload.mnemonic_keyword.strip(),
            "verbal_cue": payload.mnemonic_cue.strip()
        }

    card = Card(
        phrase_id=phrase.id,
        user_id=current_user,
        subject=clean_sub,
        text=payload.text.strip(),
        secondary_text=payload.secondary_text.strip(),
        translation=payload.translation.strip(),
        example=payload.example.strip(),
        difficulty=payload.difficulty,
        stability=1.5 if mnemonic_json else 1.0,
        state=0,
        mnemonic=mnemonic_json,
        next_review=datetime.utcnow()
    )
    db.add(card)
    await db.flush()
    saved_id = card.id
    await db.commit()
    return {"status": "success", "card_id": saved_id}

@router.put("/management/cards/{card_id}")
async def update_single_card(
    card_id: int,
    payload: CardUpdateIn,
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    card_res = await db.execute(select(Card).filter(Card.id == card_id, Card.user_id == current_user))
    card = card_res.scalar_one_or_none()
    if not card:
        raise HTTPException(status_code=404, detail="Карточка не найдена или нет прав доступа")

    card.text = payload.text.strip()
    card.secondary_text = payload.secondary_text.strip()
    card.translation = payload.translation.strip()
    card.example = payload.example.strip()

    if payload.mnemonic_keyword or payload.mnemonic_cue:
        card.mnemonic = {
            "keyword": payload.mnemonic_keyword.strip(),
            "verbal_cue": payload.mnemonic_cue.strip()
        }

    await db.commit()
    return {"status": "success", "card_id": card.id}

# --- 6. МИГРАЦИЯ КАРТОЧЕК МЕЖДУ ПРЕДМЕТАМИ ---
@router.post("/management/cards/{card_id}/move")
async def move_card(
    card_id: int, 
    payload: CardMoveIn, 
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    card_res = await db.execute(select(Card).filter(Card.id == card_id, Card.user_id == current_user))
    card = card_res.scalar_one_or_none()
    if not card:
        raise HTTPException(status_code=404, detail="Карточка не найдена или нет прав доступа")
    
    target_sub = payload.target_subject.strip().lower()
    
    phrase_res = await db.execute(
        select(Phrase).filter(Phrase.text == "[МИГРИРОВАВШИЕ КАРТОЧКИ]", Phrase.subject == target_sub, Phrase.user_id == current_user)
    )
    phrase = phrase_res.scalar_one_or_none()
    if not phrase:
        phrase = Phrase(text="[МИГРИРОВАВШИЕ КАРТОЧКИ]", subject=target_sub, user_id=current_user)
        db.add(phrase)
        await db.flush()
        
    card.subject = target_sub
    card.phrase_id = phrase.id
    await db.commit()
    
    return {"status": "success", "card_id": card_id, "target_subject": target_sub}

# --- 7. БЕЗОПАСНОЕ УДАЛЕНИЕ КАРТОЧЕК ---
@router.delete("/management/cards/{card_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_card(
    card_id: int, 
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    card_res = await db.execute(select(Card).filter(Card.id == card_id, Card.user_id == current_user))
    card = card_res.scalar_one_or_none()
    if not card: 
        raise HTTPException(status_code=404, detail="Карточка не найдена или нет прав доступа")
    await db.delete(card)
    await db.commit()
    return None

# --- 7.5 ПЕРЕГЕНЕРАЦИЯ АССОЦИАЦИИ ---
@router.post("/management/cards/{card_id}/regenerate_mnemonic")
async def regenerate_mnemonic(
    card_id: int, 
    payload: RegenerateMnemonicIn, 
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    card_res = await db.execute(select(Card).filter(Card.id == card_id, Card.user_id == current_user))
    card = card_res.scalar_one_or_none()
    if not card: 
        raise HTTPException(status_code=404, detail="Карточка не найдена или нет прав доступа")
        
    new_mnemonic = await regenerate_card_mnemonic(
        text=card.text, 
        translation=card.translation, 
        subject=card.subject, 
        preference=payload.preference
    )
    
    if "error" in new_mnemonic:
        raise HTTPException(status_code=500, detail=new_mnemonic["error"])
        
    card.mnemonic = new_mnemonic
    await db.commit()
    return {"status": "success", "mnemonic": new_mnemonic}

# --- 8. МАССОВЫЙ ПЕРЕНОС КАРТОЧЕК ---
@router.post("/data/cards/move")
async def bulk_move_cards(
    payload: BulkCardMoveIn, 
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    if not payload.card_ids or not payload.target_subject.strip():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Идентификаторы карточек не могут быть пустыми и целевой предмет должен быть указан"
        )
    
    target_sub = payload.target_subject.strip().lower()
    
    phrase_res = await db.execute(
        select(Phrase).filter(Phrase.text == "[МИГРИРОВАВШИЕ КАРТОЧКИ]", Phrase.subject == target_sub, Phrase.user_id == current_user)
    )
    phrase = phrase_res.scalar_one_or_none()
    if not phrase:
        phrase = Phrase(text="[МИГРИРОВАВШИЕ КАРТОЧКИ]", subject=target_sub, user_id=current_user)
        db.add(phrase)
        await db.flush()
    
    stmt = (
        update(Card)
        .where(Card.id.in_(payload.card_ids), Card.user_id == current_user)
        .values(subject=target_sub, phrase_id=phrase.id)
    )
    await db.execute(stmt)
    await db.commit()
    
    return {
        "status": "success",
        "moved_count": len(payload.card_ids),
        "target_subject": target_sub
    }

# --- 8.5 МАССОВОЕ УДАЛЕНИЕ КАРТОЧЕК И ОЧИСТКА ПРЕДМЕТА ---
@router.post("/data/cards/delete")
async def bulk_delete_cards(
    payload: BulkCardDeleteIn,
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    if not payload.card_ids:
        raise HTTPException(status_code=400, detail="Список идентификаторов пуст")
    await db.execute(delete(Card).where(Card.id.in_(payload.card_ids), Card.user_id == current_user))
    await db.commit()
    return {"status": "success", "deleted_count": len(payload.card_ids)}

# --- 8.1 УПРАВЛЕНИЕ ПРЕДМЕТАМИ (СПИСОК, ПЕРЕИМЕНОВАНИЕ, УДАЛЕНИЕ) ---
@router.get("/data/subjects/details")
async def get_subjects_details(
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    stmt = (
        select(Card.subject, func.count(Card.id))
        .filter(Card.user_id == current_user)
        .group_by(Card.subject)
    )
    res = await db.execute(stmt)
    rows = res.all()

    phrase_stmt = select(Phrase.subject).filter(Phrase.user_id == current_user).distinct()
    phrase_res = await db.execute(phrase_stmt)
    phrase_subs = [s[0] for s in phrase_res.all() if s[0]]

    counts_map = {sub: count for sub, count in rows if sub}
    for ps in phrase_subs:
        if ps not in counts_map:
            counts_map[ps] = 0

    subjects_list = [
        {"slug": sub, "name": sub.upper(), "cards_count": counts_map[sub]}
        for sub in sorted(counts_map.keys())
    ]
    return {"status": "success", "subjects": subjects_list}

@router.post("/data/subjects/rename")
async def rename_subject(
    payload: SubjectRenameIn,
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    old_sub = payload.old_subject.strip().lower()
    new_sub = payload.new_subject.strip().lower()
    
    if not old_sub or not new_sub:
        raise HTTPException(status_code=400, detail="Название предмета не может быть пустым.")
    if old_sub == "all" or new_sub == "all":
        raise HTTPException(status_code=400, detail="Нельзя использовать зарезервированное имя 'all'.")
    if old_sub == new_sub:
        return {"status": "success", "message": "Имена совпадают", "subject": new_sub, "cards_updated": 0}

    # 1. Обновляем карточки
    card_res = await db.execute(
        update(Card)
        .where(Card.subject == old_sub, Card.user_id == current_user)
        .values(subject=new_sub)
    )
    cards_updated = card_res.rowcount

    # 2. Обновляем темы (Phrase)
    await db.execute(
        update(Phrase)
        .where(Phrase.subject == old_sub, Phrase.user_id == current_user)
        .values(subject=new_sub)
    )

    # 3. Обновляем задачи генерации (GenerationJob)
    await db.execute(
        update(GenerationJob)
        .where(GenerationJob.subject == old_sub, GenerationJob.user_id == current_user)
        .values(subject=new_sub)
    )

    # 4. Обновляем лимиты в UserSetting
    setting_res = await db.execute(select(UserSetting).filter(UserSetting.user_id == current_user))
    setting = setting_res.scalar_one_or_none()
    if setting and setting.subject_limits and old_sub in setting.subject_limits:
        limits = dict(setting.subject_limits)
        limits[new_sub] = limits.pop(old_sub)
        setting.subject_limits = limits

    await db.commit()
    return {
        "status": "success",
        "old_subject": old_sub,
        "new_subject": new_sub,
        "cards_updated": cards_updated
    }

@router.delete("/data/subjects/{subject_slug}")
async def delete_subject_all(
    subject_slug: str,
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    await check_experiment_lock(current_user, db)
    sub = subject_slug.strip().lower()
    if sub == "all":
        raise HTTPException(status_code=400, detail="Нельзя удалить служебный фильтр 'all'.")

    # Удаляем ReviewLog карточек предмета, чтобы не оставалось повисших записей
    card_ids_res = await db.execute(select(Card.id).where(Card.subject == sub, Card.user_id == current_user))
    card_ids = card_ids_res.scalars().all()
    if card_ids:
        await db.execute(delete(ReviewLog).where(ReviewLog.card_id.in_(card_ids)))
        await db.execute(delete(Card).where(Card.id.in_(card_ids)))

    await db.execute(delete(Phrase).where(Phrase.subject == sub, Phrase.user_id == current_user))
    await db.execute(delete(GenerationJob).where(GenerationJob.subject == sub, GenerationJob.user_id == current_user))

    # Очищаем лимиты из UserSetting
    setting_res = await db.execute(select(UserSetting).filter(UserSetting.user_id == current_user))
    setting = setting_res.scalar_one_or_none()
    if setting and setting.subject_limits and sub in setting.subject_limits:
        limits = dict(setting.subject_limits)
        del limits[sub]
        setting.subject_limits = limits

    await db.commit()
    return {"status": "success", "deleted_subject": sub, "deleted_cards": len(card_ids)}

# --- 9. ТАЙМЕР ПОМОДОРО ---
@router.post("/timer/rest")
async def start_rest_session(
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    now = datetime.utcnow()
    rest_end = now + timedelta(minutes=17)
    session_res = await db.execute(select(UserSession).filter(UserSession.telegram_id == current_user))
    session = session_res.scalar_one_or_none()
    if not session: 
        session = UserSession(telegram_id=current_user, user_id=current_user)
        db.add(session)
    session.is_resting = True
    session.rest_ends_at = rest_end
    session.notified = False
    await db.commit()
    return {"status": "rest_started", "rest_ends_at": rest_end}

@router.get("/timer/status")
async def get_timer_status(
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    session_res = await db.execute(select(UserSession).filter(UserSession.telegram_id == current_user))
    session = session_res.scalar_one_or_none()
    if not session or not session.is_resting: 
        return {"is_resting": False, "seconds_left": 0}
    now = datetime.utcnow()
    if now >= session.rest_ends_at:
        session.is_resting = False
        await db.commit()
        return {"is_resting": False, "seconds_left": 0}
    return {"is_resting": True, "seconds_left": int((session.rest_ends_at - now).total_seconds())}

# --- 10. ЗАПИСЬ ЕЖЕДНЕВНЫХ ИТОГОВ И КОГНИТИВНОГО ОПРОСА ---
@router.post("/stats/daily_session")
async def log_daily_session(
    payload: DailySessionIn, 
    current_user: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db)
):
    check_time = datetime.utcnow() - timedelta(seconds=60)
    dup_stmt = select(DailySession).filter(
        DailySession.user_id == current_user,
        DailySession.timestamp >= check_time
    )
    dup_res = await db.execute(dup_stmt)
    if dup_res.scalars().first():
        return {"status": "success", "message": "Дубликат пропущен"}

    last_session_stmt = select(DailySession).filter(DailySession.user_id == current_user).order_by(DailySession.id.desc()).limit(1)
    last_session_res = await db.execute(last_session_stmt)
    last_session = last_session_res.scalar_one_or_none()
    
    start_time = datetime.utcnow() - timedelta(hours=24)
    if last_session:
        start_time = last_session.timestamp
        
    logs_stmt = select(ReviewLog).filter(
        ReviewLog.user_id == current_user,
        ReviewLog.review_time > start_time
    )
    logs_res = await db.execute(logs_stmt)
    logs = logs_res.scalars().all()
    
    total_reviewed = len(logs)
    new_cards_learned = sum(1 for l in logs if l.state == 0)
    recalls = sum(1 for l in logs if l.rating in (2, 3, 4))
    true_retention = float(recalls) / total_reviewed if total_reviewed > 0 else 0.0
    
    count_stmt = select(func.count(DailySession.id)).filter(DailySession.user_id == current_user)
    count_res = await db.execute(count_stmt)
    total_sessions_count = count_res.scalar_one()
    date_marker = f"День {total_sessions_count + 1}"
    
    session_entry = DailySession(
        user_id=current_user,
        date=date_marker,
        total_reviewed=total_reviewed,
        new_cards_learned=new_cards_learned,
        session_duration=payload.session_duration,
        true_retention=true_retention,
        mental_effort=payload.mental_effort,
        association_utility=payload.association_utility,
        perceived_retention=payload.perceived_retention,
        timestamp=datetime.utcnow()
    )
    db.add(session_entry)
    await db.commit()
    
    return {
        "status": "success",
        "date": date_marker,
        "total_reviewed": total_reviewed,
        "new_cards_learned": new_cards_learned,
        "true_retention": true_retention
    }