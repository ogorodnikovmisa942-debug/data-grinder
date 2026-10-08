"""Текст из загруженного файла: PDF, DOCX, PPTX, TXT/MD (без ИИ). Нужен там, где файл не идёт в нарезку книги (например, список билетов)."""
import io
import re

MAX_EXTRACT_BYTES = 8 * 1024 * 1024
SUPPORTED = (".pdf", ".docx", ".pptx", ".txt", ".md")


class ExtractError(ValueError):
    pass


def extract_text(filename: str, contents: bytes) -> str:
    name = (filename or "").lower()
    if len(contents) > MAX_EXTRACT_BYTES:
        raise ExtractError("Файл слишком большой (максимум 8 МБ).")
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        pages = [(p.extract_text() or "").strip() for p in PdfReader(io.BytesIO(contents)).pages]
        text = "\n".join(p for p in pages if p)
    elif name.endswith(".docx"):
        import docx
        doc = docx.Document(io.BytesIO(contents))
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        text = "\n".join(lines)
    elif name.endswith(".pptx"):
        from pptx import Presentation
        lines = []
        for slide in Presentation(io.BytesIO(contents)).slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(shape.text.strip())
        text = "\n".join(lines)
    elif name.endswith((".txt", ".md")):
        text = contents.decode("utf-8-sig", errors="ignore")
    else:
        raise ExtractError("Этот формат не поддерживается. Подойдут PDF, Word (DOCX), PowerPoint (PPTX) и текст.")
    text = re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n"))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ExtractError("В файле не нашёлся текст.")
    return text
