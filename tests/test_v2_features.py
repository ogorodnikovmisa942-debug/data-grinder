# tests/test_v2_features.py
import io
import json
import unittest
import asyncio
from datetime import datetime, timedelta
from fastapi.testclient import TestClient
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import delete

from main import app
from app.database.session import AsyncSessionLocal
from app.database.models import Card, Phrase, ReviewLog, GenerationJob
from app.services.ai_gateway import unpack_minified_cards, split_text_into_chunks

class TestV2Features(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client_cm = TestClient(app)
        cls.client = cls.client_cm.__enter__()
        cls.user_id = "test_v2_user"

    @classmethod
    def tearDownClass(cls):
        cls.client_cm.__exit__(None, None, None)

    def run_async(self, coro):
        return asyncio.run(coro)

    def test_01_cloze_unpacking(self):
        """Проверка автораспознавания карточек с пропусками {{c1::ответ::подсказка}}."""
        raw_payload = {
            "domain": "law",
            "slug": "law_const",
            "title": "Конституционное право",
            "c": [
                {
                    "t": "Судебная власть в РБ принадлежит исключительно {{c1::судам::орган}}.",
                    "s": "Конституция | ст. 109",
                    "d": "Судам.",
                    "e": "Правосудие осуществляется только судом.",
                    "l": "easy",
                    "h": "Судебная власть"
                },
                {
                    "t": "Что такое норма права?",
                    "s": "Теория права",
                    "d": "Общеобязательное правило поведения.",
                    "e": "Норма закреплена в законе.",
                    "l": "easy",
                    "h": "Общая теория"
                }
            ]
        }
        unpacked = unpack_minified_cards(raw_payload)
        cards = unpacked["cards"]
        self.assertEqual(len(cards), 2)
        self.assertEqual(cards[0]["content_type"], "cloze")
        self.assertIn("{{c1::судам::орган}}", cards[0]["text"])
        self.assertEqual(cards[1]["content_type"], "text")

    def test_02_chunking_economics_and_overlap(self):
        """Проверка деления текста на чанки 14 000 символов с overlap 1000 символов."""
        sample_text = "Параграф текста для тестирования нарезки знаний. " * 500  # ~24 500 символов
        chunks = split_text_into_chunks(sample_text, max_chunk_chars=14000, overlap_chars=1000)
        self.assertGreaterEqual(len(chunks), 2)
        for ch in chunks:
            self.assertLessEqual(len(ch), 15000)
            self.assertGreater(len(ch), 0)

    def test_03_docx_file_import(self):
        """Проверка извлечения текста из реального файла Microsoft Word (.docx)."""
        doc = Document()
        doc.add_heading("Тестовая лекция по биохимии", level=1)
        doc.add_paragraph("Митохондрии являются энергетическими станциями клетки.")
        doc.add_paragraph("Синтез АТФ происходит на внутренней мембране митохондрий.")
        
        docx_io = io.BytesIO()
        doc.save(docx_io)
        docx_io.seek(0)

        response = self.client.post(
            "/api/config/import/file",
            headers={"X-User-Id": self.user_id},
            data={"subject": "biochemistry", "is_deferred": "true"},
            files={"files": ("lecture.docx", docx_io, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "queued")
        
        # Проверяем, что в БД создан GenerationJob с корректно извлеченным текстом
        async def verify_job():
            async with AsyncSessionLocal() as db:
                job = await db.get(GenerationJob, data["job_id"])
                return job.raw_text

        raw_text = self.run_async(verify_job())
        self.assertIn("Митохондрии являются энергетическими станциями клетки", raw_text)
        self.assertIn("Синтез АТФ", raw_text)

    def test_04_pptx_file_import(self):
        """Проверка извлечения текста из реальной презентации Microsoft PowerPoint (.pptx)."""
        prs = Presentation()
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        tx_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        tf = tx_box.text_frame
        p = tf.paragraphs[0]
        p.text = "Слайд 1: Принципы гражданского процесса"

        pptx_io = io.BytesIO()
        prs.save(pptx_io)
        pptx_io.seek(0)

        response = self.client.post(
            "/api/config/import/file",
            headers={"X-User-Id": self.user_id},
            data={"subject": "civil_procedure", "is_deferred": "true"},
            files={"files": ("slides.pptx", pptx_io, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "queued")

        async def verify_job():
            async with AsyncSessionLocal() as db:
                job = await db.get(GenerationJob, data["job_id"])
                return job.raw_text

        raw_text = self.run_async(verify_job())
        self.assertIn("Принципы гражданского процесса", raw_text)

    def test_05_deck_sharing_and_analytics(self):
        """Проверка создания шеринговой ссылки колоды и получения зрелости/тепловой карты."""
        async def setup_cards():
            async with AsyncSessionLocal() as db:
                # Очищаем старые тестовые данные
                await db.execute(delete(Card).filter(Card.user_id == self.user_id))
                await db.execute(delete(Phrase).filter(Phrase.user_id == self.user_id))
                await db.commit()

                phrase = Phrase(text="Шеринг тест блок", subject="share_test_sub", user_id=self.user_id)
                db.add(phrase)
                await db.flush()

                c1 = Card(
                    phrase_id=phrase.id,
                    user_id=self.user_id,
                    subject="share_test_sub",
                    text="Тест вопрос 1",
                    secondary_text="Подсказка 1",
                    translation="Ответ 1",
                    state=2, # Review
                    stability=45.0, # Mature
                    difficulty=3.0,
                    content_type="text",
                    next_review=datetime.utcnow() - timedelta(hours=1)
                )
                c2 = Card(
                    phrase_id=phrase.id,
                    user_id=self.user_id,
                    subject="share_test_sub",
                    text="Тест вопрос 2 с пропуском {{c1::ответ::подсказка}}",
                    secondary_text="Подсказка 2",
                    translation="Ответ 2",
                    state=0, # New
                    stability=0.0,
                    difficulty=5.0,
                    content_type="cloze",
                    next_review=datetime.utcnow()
                )
                db.add_all([c1, c2])
                await db.commit()
                return phrase.id

        phrase_id = self.run_async(setup_cards())

        # 2. Создаем ссылку на колоду
        share_res = self.client.post(
            "/api/data/cards/share",
            headers={"X-User-Id": self.user_id},
            json={"subject": "share_test_sub", "title": "Тестовая колода для друзей"}
        )
        self.assertEqual(share_res.status_code, 200)
        share_data = share_res.json()
        self.assertEqual(share_data["status"], "success")
        self.assertIn("share_url", share_data)
        self.assertIn("share_key", share_data)
        self.assertEqual(share_data["total_cards"], 2)

        # 3. Запрашиваем дашборд аналитики
        stats_res = self.client.get(
            f"/api/stats/dashboard?subject=share_test_sub",
            headers={"X-User-Id": self.user_id}
        )
        self.assertEqual(stats_res.status_code, 200)
        stats = stats_res.json()
        self.assertIn("maturity", stats)
        self.assertIn("heatmap", stats)
        self.assertEqual(stats["maturity"]["new"], 1)
        self.assertEqual(stats["maturity"]["mature"], 1)
        self.assertEqual(stats["total_cards"], 2)

    def test_06_sibling_burying_for_cloze_cards(self):
        """Проверка Sibling Burying: если sibling изучен сегодня, второй скрывается до завтра."""
        async def setup_siblings():
            async with AsyncSessionLocal() as db:
                sub = "sibling_test_sub"
                await db.execute(delete(Card).filter(Card.user_id == "user_sibling"))
                await db.execute(delete(Phrase).filter(Phrase.user_id == "user_sibling"))
                await db.execute(delete(ReviewLog).filter(ReviewLog.user_id == "user_sibling"))
                await db.commit()

                phrase = Phrase(text="Семья клоз карт", subject=sub, user_id="user_sibling")
                db.add(phrase)
                await db.flush()

                c_sib1 = Card(
                    phrase_id=phrase.id,
                    user_id="user_sibling",
                    subject=sub,
                    text="В 1945 году {{c1::ООН}} была основана в Сан-Франциско.",
                    translation="ООН",
                    state=2,
                    stability=5.0,
                    content_type="cloze",
                    next_review=datetime.utcnow() - timedelta(hours=2)
                )
                c_sib2 = Card(
                    phrase_id=phrase.id,
                    user_id="user_sibling",
                    subject=sub,
                    text="В {{c1::1945}} году ООН была основана в Сан-Франциско.",
                    translation="1945",
                    state=2,
                    stability=5.0,
                    content_type="cloze",
                    next_review=datetime.utcnow() - timedelta(hours=2)
                )
                db.add_all([c_sib1, c_sib2])
                await db.flush()

                # Симулируем, что c_sib1 уже был повторен сегодня
                log = ReviewLog(
                    card_id=c_sib1.id,
                    user_id="user_sibling",
                    rating=3,
                    state=2,
                    review_time=datetime.utcnow()
                )
                db.add(log)
                await db.commit()
                return c_sib1.id, c_sib2.id

        sib1_id, sib2_id = self.run_async(setup_siblings())

        # Запрашиваем активную сессию обучения
        res = self.client.get(
            "/api/session?subject=sibling_test_sub",
            headers={"X-User-Id": "user_sibling"}
        )
        self.assertEqual(res.status_code, 200)
        session_cards = res.json()
        card_ids = [c["id"] for c in session_cards]
        # c_sib2 должен быть изолирован (захоронен / buried), так как c_sib1 повторялся сегодня
        self.assertNotIn(sib2_id, card_ids)
