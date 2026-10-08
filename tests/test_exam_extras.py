"""Экзамен: повторный разбор недостающих билетов, готовность и прогноз, аварийный режим, симулятор, приоритеты предметов, напоминания;
экспорт в Anki и чтение файла со списком билетов."""
import asyncio
import io
import sqlite3
import tempfile
import zipfile
from datetime import date, timedelta
from types import SimpleNamespace
from unittest.mock import patch

from sqlalchemy import select

import test_exam_prep as tep
from app.database.models import Card, ExamPlan, ExamTicket
from app.database.session import AsyncSessionLocal
from app.services import exam_prep as ep
from app.services.anki_export import build_apkg, deck_id_for
from app.services.knowledge_path import complete_lesson
from app.services.text_extract import ExtractError, extract_text


# --- чистая арифметика ---------------------------------------------------------------------------------------------------

def _card(state, stability=0.0):
    return SimpleNamespace(state=state, stability=stability)


def test_ticket_strength_grows_from_untouched_to_locked_in():
    assert ep.ticket_strength(None, 0.0, False, 10) == 0.0
    assert ep.ticket_strength(None, 0.5, False, 10) == 0.05                  # часть тем пройдена
    assert ep.ticket_strength(_card(0), 1.0, True, 10) == 0.15                # темы пройдены, билет не отвечен
    assert ep.ticket_strength(_card(1), 1.0, True, 10) == 0.4                 # заучивается
    assert ep.ticket_strength(_card(2, stability=3), 1.0, True, 10) == 0.7    # повторение, но стойкость короче срока
    assert ep.ticket_strength(_card(2, stability=30), 1.0, True, 10) == 1.0   # запомнится до экзамена


def _t(i, status="ok", ids=()):
    return SimpleNamespace(id=i, order_idx=i, question=f"Билет {i}", node_ids=list(ids), status=status)


def test_greedy_picks_the_tickets_that_cost_the_fewest_new_topics():
    a, b, c, d = _t(1), _t(2), _t(3), _t(4)
    needs = {1: frozenset({1}), 2: frozenset({2}), 3: frozenset({3, 4, 5}), 4: frozenset({1, 2})}
    chosen, covered = ep.greedy_tickets([a, b, c, d], needs, 3)
    assert [t.id for t in chosen] == [1, 2, 4] and covered == {1, 2}           # билет 4 закрывается уже взятыми темами
    chosen, covered = ep.greedy_tickets([a, b, c, d], needs, 10)
    assert {t.id for t in chosen} == {1, 2, 3, 4} and covered == {1, 2, 3, 4, 5}
    assert ep.greedy_tickets([c], {3: frozenset({3, 4, 5})}, 2) == ([], set())


def _ctx(tickets, needs_by_ticket, studied=(), days_left=5, remaining=None):
    nodes = [{"id": i, "key": f"n{i}", "prereq_keys": []} for i in range(1, 8)]
    for t in tickets:
        t.node_ids = sorted(needs_by_ticket[t.id])
    need_all = set().union(*needs_by_ticket.values()) if needs_by_ticket else set()
    return {"active": tickets, "tickets": tickets, "nodes": nodes, "studied": set(studied), "days_left": days_left,
            "remaining": len(need_all - set(studied)) if remaining is None else remaining, "quota": 2, "cards": {}}


def test_emergency_is_needed_only_when_the_topics_do_not_fit_and_keeps_the_cheapest_tickets():
    tickets = [_t(1), _t(2), _t(3), _t(4)]
    ctx = _ctx(tickets, {1: {1}, 2: {2}, 3: {3, 4, 5}, 4: {1, 2}}, days_left=3)       # 1 учебный день × 3 темы
    info = ep._emergency(ctx)
    assert info["needed"] and info["capacity"] == 3 and info["keep"] == 3
    assert [p["id"] for p in info["postpone"]] == [3] and info["lessons_after"] == 2
    roomy = ep._emergency(_ctx(tickets, {1: {1}, 2: {2}, 3: {3, 4, 5}, 4: {1, 2}}, days_left=10))
    assert not roomy["needed"] and roomy["postpone"] == []                              # времени хватает: откладывать нечего
    past = ep._emergency(_ctx(tickets, {1: {1}, 2: {2}, 3: {3}, 4: {4}}, days_left=-1))
    assert not past["needed"]


def test_forecast_uses_the_real_pace_or_the_plan_when_there_is_no_history():
    tickets = [_t(1), _t(2), _t(3)]
    ctx = _ctx(tickets, {1: {1}, 2: {2}, 3: {3, 4}}, days_left=6)                       # 4 учебных дня
    fc = ep.exam_forecast(ctx, pace=0.5)                                                 # 2 темы: закроются билеты 1 и 2
    assert fc["capacity"] == 2 and fc["tickets_closable"] == 2 and fc["will_make_it"] is False
    fc = ep.exam_forecast(ctx, pace=0.0)                                                 # истории нет: считаем по плану (quota=2 → 8 тем)
    assert fc["capacity"] == 8 and fc["will_make_it"] is True and fc["tickets_closable"] == 3


def test_reminder_text_counts_days_and_stays_quiet_outside_the_window():
    info = {"subject": "право", "exam_date": date(2026, 10, 20), "tickets": 40, "not_strong": 14}
    assert ep.exam_reminder_text(info, date(2026, 10, 14)) == "До экзамена 6 дней по предмету «право»: 14 из 40 билетов не закреплено."
    assert ep.exam_reminder_text(info, date(2026, 10, 19)).startswith("До экзамена 1 день ")
    assert ep.exam_reminder_text(info, date(2026, 10, 20)).startswith("Сегодня экзамен ")
    assert ep.exam_reminder_text(info, date(2026, 10, 21)) is None                       # экзамен прошёл
    assert ep.exam_reminder_text(info, date(2026, 6, 1)) is None                         # рано
    assert "все билеты закреплены" in ep.exam_reminder_text({**info, "not_strong": 0}, date(2026, 10, 14))
    assert ep.exam_reminder_text({**info, "tickets": 0}, date(2026, 10, 14)) is None
    assert [ep._days_word(n) for n in (1, 2, 5, 11, 12, 21, 22, 25)] == ["день", "дня", "дней", "дней", "дней", "день", "дня", "дней"]


# --- сценарии с базой ---------------------------------------------------------------------------------------------------

async def _make_plan(days: int = 10):
    nodes = await tep._build()
    async with AsyncSessionLocal() as db:
        plan = await ep.create_plan(db, tep.USER, tep.SUBJECT, "Билеты", date.today() + timedelta(days=days), [
            {"question": "Расскажите о теме", "answer": ""}, {"question": "Вопрос не из книги", "answer": ""}])
        plan_id = plan.id
    with patch("app.services.ai_gateway.exam_matcher.match_tickets", side_effect=tep._fake_match):
        await ep.run_plan_matching(plan_id)
    return nodes, plan_id


def test_rematch_looks_only_at_missing_tickets_and_keeps_the_rest():
    seen = {}

    async def second_match(nodes, cards_by_node, questions, calls):
        seen["questions"] = list(questions)
        calls.append({"label": "exam#2.1", "cost_usd": 0.002, "prompt_tokens": 1, "cache_hit_tokens": 0, "completion_tokens": 1,
                      "duration_ms": 1, "finish_reason": "stop", "model": "deepseek-flash"})
        return [{"nodes": ["other"], "found": True, "points": [{"text": "Другая тема", "variants": [], "weight": 1.0}]}]

    async def scenario():
        nodes, plan_id = await _make_plan()
        try:
            with patch("app.services.ai_gateway.exam_matcher.match_tickets", side_effect=second_match):
                closed = await ep.run_plan_matching(plan_id, only_missing=True)
            assert closed == 1 and seen["questions"] == ["Вопрос не из книги"]            # первый билет второй раз не отправляли
            async with AsyncSessionLocal() as db:
                tickets = (await db.execute(select(ExamTicket).where(ExamTicket.plan_id == plan_id).order_by(ExamTicket.order_idx))).scalars().all()
                assert [t.status for t in tickets] == ["ok", "ok"] and tickets[1].node_ids == [nodes["other"]] and tickets[1].card_id
                plan = (await db.execute(select(ExamPlan).where(ExamPlan.id == plan_id))).scalar_one()
                assert plan.status == "ready" and round(plan.cost_usd, 4) == 0.006       # 0.004 + 0.002
            assert await ep.run_plan_matching(plan_id, only_missing=True) == 0            # недостающих больше нет — ничего не делаем
        finally:
            await tep._cleanup()

    asyncio.run(scenario())


def test_rematch_failure_leaves_the_plan_ready():
    async def broken(nodes, cards_by_node, questions, calls):
        raise RuntimeError("DeepSeek недоступен")

    async def scenario():
        nodes, plan_id = await _make_plan()
        try:
            with patch("app.services.ai_gateway.exam_matcher.match_tickets", side_effect=broken):
                assert await ep.run_plan_matching(plan_id, only_missing=True) == 0
            async with AsyncSessionLocal() as db:
                plan = (await db.execute(select(ExamPlan).where(ExamPlan.id == plan_id))).scalar_one()
                assert plan.status == "ready"                                              # сбой не ломает готовый план
        finally:
            await tep._cleanup()

    asyncio.run(scenario())


def test_overview_shows_readiness_and_the_simulator_draws_only_studied_tickets():
    async def scenario():
        nodes, plan_id = await _make_plan()
        try:
            async with AsyncSessionLocal() as db:
                ov = await ep.exam_overview(db, tep.USER, tep.SUBJECT)
                assert ov["readiness"]["percent"] == 0 and ov["readiness"]["tickets_total"] == 1
                assert ov["emergency"]["needed"] is False and ov["tickets"]["postponed"] == 0
                try:
                    await ep.simulator_draw(db, tep.USER, tep.SUBJECT)
                    assert False, "темы не пройдены: билета для симулятора быть не должно"
                except LookupError:
                    pass
                for key in ("base", "topic"):
                    await complete_lesson(db, tep.USER, nodes[key], 0)
                await db.commit()
                await tep._learn(db, nodes["topic"])
                draw = await ep.simulator_draw(db, tep.USER, tep.SUBJECT)
                assert draw["question"] == "Расскажите о теме" and draw["seconds"] == ep.SIMULATOR_SECONDS and draw["points"] == 1
                good = await ep.simulator_check(db, tep.USER, draw["ticket_id"], "Тема — это главное в курсе")
                assert good["percent"] == 100 and good["hit"] == ["Тема — главное"] and good["missed"] == []
                bad = await ep.simulator_check(db, tep.USER, draw["ticket_id"], "Не помню")
                assert bad["percent"] == 0 and bad["missed"] == ["Тема — главное"] and "главное" in bad["reference"]
                assert (await ep.exam_overview(db, tep.USER, tep.SUBJECT))["readiness"]["percent"] > 0
                pr = await ep.exam_priorities(db, tep.USER)
                assert [p["subject"] for p in pr] == [tep.SUBJECT] and 0 < pr[0]["priority"] <= 1
                rem = await ep.exam_reminders(db, [tep.USER])
                assert rem[tep.USER][0]["tickets"] == 1 and rem[tep.USER][0]["not_strong"] == 1     # билет ещё не отвечен, закреплены только карточки темы
        finally:
            await tep._cleanup()

    asyncio.run(scenario())


def test_exam_api_extras():
    from fastapi.testclient import TestClient
    from main import app

    async def setup():
        return await _make_plan(days=2)

    nodes, plan_id = asyncio.run(setup())
    try:
        with patch("app.services.generation_worker.claim_next_pending_job", return_value=None):
            client = TestClient(app)
            h = {"X-User-Id": tep.USER}
            sub = tep.SUBJECT
            ov = client.get(f"/api/exam/{sub}", headers=h).json()
            assert ov["readiness"]["percent"] == 0 and "forecast" not in ov and ov["emergency"]["needed"] is True       # 2 дня: темы не успеть
            assert client.get(f"/api/exam/{sub}/simulator/draw", headers=h).status_code == 409
            em = client.get(f"/api/exam/{sub}/emergency", headers=h).json()
            assert em["available"] and em["needed"] and len(em["postpone"]) == 1
            assert client.post(f"/api/exam/{sub}/emergency/apply", headers=h).json() == {"postponed": 1}
            ov = client.get(f"/api/exam/{sub}", headers=h).json()
            assert ov["tickets"]["postponed"] == 1 and ov["tickets"]["ok"] == 0
            assert client.post(f"/api/exam/{sub}/emergency/undo", headers=h).json() == {"restored": 1}
            assert client.get(f"/api/exam/{sub}", headers=h).json()["tickets"]["ok"] == 1
            assert client.post(f"/api/exam/{sub}/rematch", headers=h).status_code == 200                            # есть билет «нет в книге»
            pr = client.get("/api/exams/priorities", headers=h).json()["items"]
            assert pr and pr[0]["subject"] == sub
            up = client.post("/api/exam/extract", headers=h, files={"file": ("tickets.txt", "1. Понятие права\n2. Функции государства\n".encode("utf-8"))})
            assert up.status_code == 200 and up.json()["count"] == 2 and "Понятие права" in up.json()["text"]
            bad = client.post("/api/exam/extract", headers=h, files={"file": ("tickets.exe", b"zzz")})
            assert bad.status_code == 400
            assert client.get(f"/api/exam/{sub}/emergency", headers={"X-User-Id": "someone_else"}).json() == {"available": False}
    finally:
        asyncio.run(tep._cleanup())


# --- чтение файла со списком билетов -------------------------------------------------------------------------------------

def test_extract_text_reads_text_and_word_files_and_rejects_the_rest():
    assert extract_text("a.txt", "1. Первый\r\n\r\n\r\n\r\n2. Второй".encode("utf-8")) == "1. Первый\n\n2. Второй"
    import docx
    d = docx.Document()
    d.add_paragraph("1. Понятие государства")
    d.add_paragraph("2. Функции государства")
    buf = io.BytesIO()
    d.save(buf)
    assert extract_text("список.DOCX", buf.getvalue()) == "1. Понятие государства\n2. Функции государства"
    for name, data in (("a.exe", b"x"), ("empty.txt", b"   ")):
        try:
            extract_text(name, data)
            assert False
        except ExtractError:
            pass


# --- Anki ----------------------------------------------------------------------------------------------------------------

def test_anki_export_builds_a_valid_apkg_with_escaped_text_and_stable_ids():
    cards = [{"id": 1, "text": "Что такое <b>право</b>?", "secondary_text": "Право | Тема", "translation": "Система норм.\nВторая строка.",
              "example": "Пример & случай", "tag": "Норма права"},
             {"id": 2, "text": "Срок полномочий?", "secondary_text": "", "translation": "Шесть лет.", "example": "", "tag": ""}]
    data = build_apkg("Data Grinder: тест", cards)
    assert data[:2] == b"PK" and build_apkg("Data Grinder: тест", cards) != b"" and deck_id_for("a") == deck_id_for("a") != deck_id_for("b")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert "collection.anki2" in z.namelist() and "media" in z.namelist()
        with tempfile.NamedTemporaryFile(suffix=".anki2", delete=False) as f:
            f.write(z.read("collection.anki2"))
            path = f.name
    con = sqlite3.connect(path)
    notes = con.execute("select flds, tags, guid from notes order by id").fetchall()
    con.close()
    assert len(notes) == 2
    assert "&lt;b&gt;право&lt;/b&gt;" in notes[0][0] and "Система норм.<br>Вторая строка." in notes[0][0] and "Пример &amp; случай" in notes[0][0]
    assert "Норма_права" in notes[0][1] and notes[0][2] != notes[1][2]
    import genanki
    assert notes[0][2] == genanki.guid_for("data-grinder", 1)                                # одна и та же карточка при повторном экспорте обновляется, а не дублируется


def test_anki_endpoint_returns_the_decks_cards_and_404_for_an_empty_subject():
    from fastapi.testclient import TestClient
    from main import app
    import test_knowledge_path as tkp

    async def setup():
        await tkp.cleanup()
        with patch("app.services.generation_worker.build_learning_path", side_effect=tkp.fake_build):
            await tkp.process_generation_job(await tkp.create_job(), is_offpeak=True)

    asyncio.run(setup())
    try:
        with patch("app.services.generation_worker.claim_next_pending_job", return_value=None):
            client = TestClient(app)
            h = {"X-User-Id": tkp.USER}
            r = client.get(f"/api/export/anki?subject={tkp.SUBJECT}", headers=h)
            assert r.status_code == 200 and r.content[:2] == b"PK" and r.headers["x-cards"] == "7"
            assert "attachment" in r.headers["content-disposition"]
            assert client.get("/api/export/anki?subject=нет_такого_предмета", headers=h).status_code == 404
    finally:
        asyncio.run(tkp.cleanup())


def test_apkg_import_round_trips_our_own_export_and_reads_plain_decks():
    from app.services.anki_export import ApkgError, read_apkg
    cards = [{"id": 1, "text": "Что такое <b>право</b>?", "secondary_text": "Право | Тема", "translation": "Система норм.\nВторая строка.",
              "example": "Пример & случай", "tag": ""}, {"id": 2, "text": "Срок?", "secondary_text": "", "translation": "Шесть лет.", "example": "", "tag": ""}]
    back = read_apkg(build_apkg("Тест", cards))
    assert [c["text"] for c in back] == ["Что такое <b>право</b>?", "Срок?"]                  # разметка из самого текста карточки возвращается как текст
    assert back[0]["translation"] == "Система норм.\nВторая строка." and back[0]["secondary_text"] == "Право | Тема" and back[0]["example"] == "Пример & случай"
    # простая колода из двух полей («Basic») и мусор
    import genanki
    basic = genanki.Model(1380120000, "Basic", fields=[{"name": "Front"}, {"name": "Back"}], templates=[{"name": "c", "qfmt": "{{Front}}", "afmt": "{{Back}}"}])
    deck = genanki.Deck(2059400110, "Basic deck")
    deck.add_note(genanki.Note(model=basic, fields=["Столица<br>Франции?", "Париж [sound:a.mp3]"]))
    deck.add_note(genanki.Note(model=basic, fields=["Пустой ответ", ""]))
    with tempfile.NamedTemporaryFile(suffix=".apkg", delete=False) as f:
        path = f.name
    genanki.Package(deck).write_to_file(path)
    plain = read_apkg(open(path, "rb").read())
    assert plain == [{"text": "Столица\nФранции?", "secondary_text": "", "translation": "Париж", "example": "", "initial_difficulty_tier": "medium", "mnemonic": None}]
    for bad in (b"not a zip", b"PK\x05\x06" + b"\x00" * 18):
        try:
            read_apkg(bad)
            assert False
        except ApkgError:
            pass


def test_apkg_file_goes_straight_into_the_subject_without_ai():
    from fastapi.testclient import TestClient
    from main import app
    import test_knowledge_path as tkp
    from app.database.models import Card
    cards = [{"id": i, "text": f"Вопрос импорта {i}?", "secondary_text": "", "translation": f"Ответ {i}.", "example": "", "tag": ""} for i in range(1, 4)]
    data = build_apkg("Импорт", cards)

    async def count():
        async with AsyncSessionLocal() as db:
            return len((await db.execute(select(Card).where(Card.user_id == tkp.USER, Card.subject == tkp.SUBJECT))).scalars().all())

    asyncio.run(tkp.cleanup())
    try:
        with patch("app.services.generation_worker.claim_next_pending_job", return_value=None):
            client = TestClient(app)
            r = client.post("/api/config/import/file", headers={"X-User-Id": tkp.USER}, data={"subject": tkp.SUBJECT},
                            files={"files": ("deck.apkg", data)})
            assert r.status_code == 200 and r.json()["status"] == "success" and r.json()["cards_count"] == 3
            assert asyncio.run(count()) == 3
            bad = client.post("/api/config/import/file", headers={"X-User-Id": tkp.USER}, data={"subject": tkp.SUBJECT},
                              files={"files": ("deck.apkg", b"junk")})
            assert bad.status_code == 400
    finally:
        asyncio.run(tkp.cleanup())
